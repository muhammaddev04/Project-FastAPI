import asyncio
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.db import get_sessionmaker
from app.core.time import utcnow
from app.modules.catalog import imports, service
from app.modules.catalog.models import Category, Price, Product
from app.modules.catalog.schemas import PriceIn
from app.modules.identity.deps import OrgContext
from app.modules.identity.models import Membership
from app.modules.subscriptions.models import Plan, Subscription
from tests.factories import add_member, auth, make_org, make_user


async def setup(session: AsyncSession) -> OrgContext:
    user = await make_user(session)
    org = await make_org(session, user)
    membership = (await session.scalars(select(Membership).where(Membership.organization_id == org.id))).one()
    await session.commit()
    return OrgContext(user, membership, org)


async def product(client: AsyncClient, context: OrgContext, **extra: object) -> dict[str, object]:
    response = await client.post(
        "/api/v1/catalog/products",
        headers=auth(context.user, context.organization),
        json={"sku": "TEA", "name": "Tea", "base_unit": "PCS", **extra},
    )
    assert response.status_code == 201, response.text
    return response.json()  # type: ignore[no-any-return]


async def test_cat_001_002_tenant_base_unit_and_idempotency(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    headers = auth(context.user, context.organization)
    payload = {"sku": "SKU-1", "name": "Milk", "base_unit": "L"}
    first = await client.post("/api/v1/catalog/products", headers=headers, json=payload)
    repeat = await client.post("/api/v1/catalog/products", headers=headers, json=payload)
    assert first.status_code == 201 and repeat.json() == first.json()
    unit = first.json()["units"][0]
    assert unit["is_base"] and Decimal(unit["coefficient"]) == 1 and unit["allow_fraction"]
    other = await make_org(session, context.user)
    await session.commit()
    assert (
        await client.get(f"/api/v1/catalog/products/{first.json()['id']}", headers=auth(context.user, other))
    ).status_code == 404
    assert (
        await client.post("/api/v1/catalog/products", headers=auth(context.user, other), json=payload)
    ).status_code == 201
    assert (
        await client.post("/api/v1/catalog/products", headers=auth(context.user, context.organization), json=payload)
    ).json()["error"]["code"] == "sku_taken"


@pytest.mark.parametrize("role", ["OWNER", "MANAGER", "OPERATOR", "WAREHOUSE", "COURIER"])
@pytest.mark.parametrize(
    "path,permission",
    [("/catalog/products", "catalog.view"), ("/pricing/price-lists", "pricing.view"), ("/imports", "import.run")],
)
async def test_permission_matrix_p04(
    client: AsyncClient, session: AsyncSession, role: str, path: str, permission: str
) -> None:
    context = await setup(session)
    member = context.user
    if role != "OWNER":
        member = await make_user(session)
        await add_member(session, context.organization, member, role)
        await session.commit()
    allowed = (
        role in {"OWNER", "MANAGER"}
        or permission == "catalog.view"
        and role in {"OPERATOR", "WAREHOUSE"}
        or permission == "pricing.view"
        and role == "OPERATOR"
    )
    result = await client.get(f"/api/v1{path}", headers=auth(member, context.organization))
    assert result.status_code == (200 if allowed else 403), result.text


async def test_store_denied_and_warehouse_no_prices(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    created = await product(client, context)
    warehouse = await make_user(session)
    await add_member(session, context.organization, warehouse, "WAREHOUSE")
    store = await make_org(session, context.user, "STORE")
    await session.commit()
    response = await client.get(
        f"/api/v1/catalog/products/{created['id']}", headers=auth(warehouse, context.organization)
    )
    assert response.status_code == 200 and response.json()["prices"] is None
    assert (await client.get("/api/v1/catalog/products", headers=auth(context.user, store))).status_code == 403


async def test_cat_005_006_limits_activation_and_blocking(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    sub = (await session.scalars(select(Subscription))).one()
    plan = await session.get(Plan, sub.plan_id)
    assert plan
    plan.max_products = 1
    await session.commit()
    await product(client, context)
    denied = await client.post(
        "/api/v1/catalog/products",
        headers=auth(context.user, context.organization),
        json={"sku": "SECOND", "name": "Second", "base_unit": "PCS"},
    )
    assert denied.status_code == 403 and denied.json()["error"]["details"]["current"] == 1
    inactive = await product(client, context, sku="OFF", is_active=False)
    assert (
        await client.post(
            f"/api/v1/catalog/products/{inactive['id']}/activate",
            headers=auth(context.user, context.organization),
            json={"version": 1},
        )
    ).status_code == 403
    sub.status = "FULL_BLOCK"
    await session.commit()
    denied = await client.patch(
        f"/api/v1/catalog/products/{inactive['id']}",
        headers=auth(context.user, context.organization),
        json={"version": 1, "name": "Changed"},
    )
    assert denied.json()["error"]["code"] == "subscription_blocked"


async def test_cat_007_category_tree_constraints(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    headers = auth(context.user, context.organization)
    root = await client.post("/api/v1/catalog/categories", headers=headers, json={"name": "Food"})
    child = await client.post(
        "/api/v1/catalog/categories", headers=headers, json={"name": "Tea", "parent_id": root.json()["id"]}
    )
    assert child.status_code == 201
    assert (
        await client.post(
            "/api/v1/catalog/categories", headers=headers, json={"name": "Third", "parent_id": child.json()["id"]}
        )
    ).status_code == 422
    assert (
        await client.patch(
            f"/api/v1/catalog/categories/{root.json()['id']}", headers=headers, json={"version": 1, "is_active": False}
        )
    ).json()["error"]["code"] == "category_not_empty"
    assert (await client.post("/api/v1/catalog/categories", headers=headers, json={"name": "food"})).status_code == 409
    tree = (await client.get("/api/v1/catalog/categories", headers=headers)).json()
    assert tree[0]["children"][0]["id"] == child.json()["id"]


async def test_units_and_immutable_trigger(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    created = await product(client, context)
    headers = auth(context.user, context.organization)
    base = created["units"][0]  # type: ignore[index]
    response = await client.post(
        f"/api/v1/catalog/products/{created['id']}/units",
        headers=headers,
        json={"code": "BOX24", "name": {"tg": "Қуттӣ", "ru": "Коробка", "en": "Box"}, "coefficient": "24"},
    )
    assert response.status_code == 201
    assert (
        await client.post(f"/api/v1/catalog/products/{created['id']}/units/{base['id']}/deactivate", headers=headers)
    ).status_code == 409
    with pytest.raises(DBAPIError):
        async with session.begin_nested():
            await session.execute(
                text("UPDATE product_units SET coefficient=25 WHERE id=:id"), {"id": response.json()["id"]}
            )
    from uuid import UUID

    default = await service.default_list(session, context.organization.id)
    await service.set_price(
        session, context, PriceIn(price_list_id=default.id, product_unit_id=UUID(base["id"]), price=Decimal(10))
    )
    # Box price is independent; no automatic coefficient-based pricing.
    assert await service.resolve(session, default.id, UUID(response.json()["id"]), utcnow()) is None


async def test_prc_001_002_003_history_future_cancel_resolve(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    created = await product(client, context)
    unit_id = created["units"][0]["id"]  # type: ignore[index]
    headers = auth(context.user, context.organization)
    lists = (await client.get("/api/v1/pricing/price-lists", headers=headers)).json()["results"]
    default_id = lists[0]["id"]
    first = await client.post(
        "/api/v1/pricing/prices",
        headers=auth(context.user, context.organization),
        json={"price_list_id": default_id, "product_unit_id": unit_id, "price": "10"},
    )
    assert first.status_code == 201, first.text
    future = await client.post(
        "/api/v1/pricing/prices",
        headers=auth(context.user, context.organization),
        json={
            "price_list_id": default_id,
            "product_unit_id": unit_id,
            "price": "12",
            "valid_from": (utcnow() + timedelta(days=2)).isoformat(),
        },
    )
    assert future.status_code == 201, future.text
    other = (
        await client.post("/api/v1/pricing/price-lists", headers=headers, json={"code": "VIP", "name": "VIP"})
    ).json()
    from uuid import UUID

    assert await service.resolve(session, UUID(other["id"]), UUID(unit_id), utcnow()) == Decimal(10)
    assert await service.resolve(session, UUID(other["id"]), UUID(unit_id), utcnow() + timedelta(days=3)) == Decimal(12)
    assert (await client.delete(f"/api/v1/pricing/prices/{future.json()['id']}", headers=headers)).status_code == 204
    await session.refresh((await session.scalars(select(Price))).one())
    assert (await session.scalars(select(Price))).one().valid_to is None
    assert (await client.delete(f"/api/v1/pricing/prices/{first.json()['id']}", headers=headers)).status_code == 409
    assert (
        await client.patch(
            f"/api/v1/pricing/price-lists/{default_id}", headers=headers, json={"version": 1, "is_active": False}
        )
    ).status_code == 409
    with pytest.raises(DBAPIError):
        async with session.begin_nested():
            await session.execute(text("UPDATE prices SET price=99 WHERE id=:id"), {"id": first.json()["id"]})


async def test_prc_concurrent_first_price_serialized(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    created = await product(client, context)
    from uuid import UUID

    default = await service.default_list(session, context.organization.id)
    await session.commit()

    async def change(amount: str) -> None:
        async with get_sessionmaker()() as other, other.begin():
            await service.set_price(
                other,
                context,
                PriceIn(
                    price_list_id=default.id, product_unit_id=UUID(created["units"][0]["id"]), price=Decimal(amount)
                ),
            )  # type: ignore[index]

    await asyncio.gather(change("10"), change("11"))
    prices = list(await session.scalars(select(Price).order_by(Price.valid_from)))
    assert len(prices) == 2 and prices[0].valid_to == prices[1].valid_from


def workbook(rows: list[list[object]], kind: str = "PRODUCTS") -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet
    sheet.append(imports.HEADERS[kind]["en"])
    for row in rows:
        sheet.append(row)
    data = BytesIO()
    book.save(data)
    return data.getvalue()


async def upload_job(client: AsyncClient, context: OrgContext, data: bytes, kind: str = "PRODUCTS") -> str:
    response = await client.post(
        "/api/v1/imports",
        headers=auth(context.user, context.organization),
        data={"kind": kind},
        files={"file": ("import.xlsx", data, imports.MIME)},
    )
    assert response.status_code == 202, response.text
    return str(response.json()["id"])


async def test_imp_products_preview_confirm_upsert_and_categories(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    await product(client, context)
    job_id = await upload_job(
        client,
        context,
        workbook(
            [["TEA", "Updated", "Drinks", None, "Good", "PCS", True], ["NEW", "New", "Drinks", None, None, "KG", True]]
        ),
    )
    assert await imports.process_pending() == 1
    headers = auth(context.user, context.organization)
    preview = (await client.get(f"/api/v1/imports/{job_id}", headers=headers)).json()
    assert preview["status"] == "VALIDATED" and preview["error_count"] == 0
    assert preview["summary"]["new_categories"] == ["Drinks"]
    assert (await client.post(f"/api/v1/imports/{job_id}/confirm", headers=headers)).status_code == 202
    await imports.process_pending()
    result = (await client.get(f"/api/v1/imports/{job_id}", headers=headers)).json()
    assert result["status"] == "COMPLETED" and result["summary"] == {"created": 1, "updated": 1, "skipped": 0}
    assert await session.scalar(select(func.count()).select_from(Category)) == 1


async def test_imp_errors_base_unit_duplicate_and_cancel(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    await product(client, context)
    job_id = await upload_job(
        client,
        context,
        workbook([["TEA", "Tea", None, None, None, "KG", True], ["TEA", "Duplicate", None, None, None, "PCS", True]]),
    )
    await imports.process_pending()
    headers = auth(context.user, context.organization)
    errors = (await client.get(f"/api/v1/imports/{job_id}/errors", headers=headers)).json()
    assert errors["count"] == 2 and errors["results"][0]["error_code"] == "base_unit_immutable"
    assert (await client.post(f"/api/v1/imports/{job_id}/confirm", headers=headers)).json()["error"][
        "code"
    ] == "import_has_errors"
    assert (await client.post(f"/api/v1/imports/{job_id}/cancel", headers=headers)).json()["status"] == "CANCELLED"
    assert (await client.post(f"/api/v1/imports/{job_id}/cancel", headers=headers)).status_code == 409


async def test_imp_all_or_nothing_limit_rollback(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    sub = (await session.scalars(select(Subscription))).one()
    plan = await session.get(Plan, sub.plan_id)
    assert plan
    plan.max_products = 1
    await session.commit()
    job_id = await upload_job(
        client,
        context,
        workbook(
            [["ONE", "One", "New category", None, None, "PCS", True], ["TWO", "Two", None, None, None, "PCS", True]]
        ),
    )
    await imports.process_pending()
    headers = auth(context.user, context.organization)
    await client.post(f"/api/v1/imports/{job_id}/confirm", headers=headers)
    await imports.process_pending()
    result = (await client.get(f"/api/v1/imports/{job_id}", headers=headers)).json()
    assert result["status"] == "FAILED" and result["failure_reason"] == "subscription_limit_reached"
    assert await session.scalar(select(func.count()).select_from(Product)) == 0
    assert await session.scalar(select(func.count()).select_from(Category)) == 0


@pytest.mark.parametrize("language", ["tg", "ru", "en"])
async def test_imp_localized_template(client: AsyncClient, session: AsyncSession, language: str) -> None:
    context = await setup(session)
    context.user.language = language
    await session.commit()
    result = await client.get("/api/v1/imports/templates/PRODUCTS", headers=auth(context.user, context.organization))
    book = load_workbook(BytesIO(result.content))
    assert book.worksheets[0]["A1"].value == imports.HEADERS["PRODUCTS"][language][0]
    assert len(book.worksheets) == 2
    book.close()


async def test_imp_005_prices_and_resolve_none(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    created = await product(client, context)
    default = await service.default_list(session, context.organization.id)
    await session.commit()
    from uuid import UUID

    unit_id = UUID(created["units"][0]["id"])  # type: ignore[index]
    assert await service.resolve(session, default.id, unit_id, utcnow()) is None
    job_id = await upload_job(client, context, workbook([["DEFAULT", "TEA", "PCS", "25.50", None]], "PRICES"), "PRICES")
    await imports.process_pending()
    headers = auth(context.user, context.organization)
    assert (await client.post(f"/api/v1/imports/{job_id}/confirm", headers=headers)).status_code == 202
    await imports.process_pending()
    assert (await client.get(f"/api/v1/imports/{job_id}", headers=headers)).json()["status"] == "COMPLETED"
    assert await service.resolve(session, default.id, unit_id, utcnow()) == Decimal("25.50")
    await client.post(
        f"/api/v1/catalog/products/{created['id']}/deactivate",
        headers=auth(context.user, context.organization),
        json={"version": 1},
    )
    assert await service.resolve(session, default.id, unit_id, utcnow()) is None


async def test_cat_003_used_base_unit_immutable(
    client: AsyncClient, session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    context = await setup(session)
    created = await product(client, context)

    class UsedReferences:
        async def has_activity(self, session: AsyncSession, product_id: object) -> bool:
            return True

    monkeypatch.setattr(service, "references", UsedReferences())
    response = await client.patch(
        f"/api/v1/catalog/products/{created['id']}",
        headers=auth(context.user, context.organization),
        json={"version": 1, "base_unit": "KG"},
    )
    assert response.status_code == 409 and response.json()["error"]["code"] == "base_unit_immutable"


async def test_price_overlap_constraint_and_cross_tenant_unit(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    created = await product(client, context)
    from uuid import UUID

    default = await service.default_list(session, context.organization.id)
    await session.commit()
    price = await service.set_price(
        session,
        context,
        PriceIn(price_list_id=default.id, product_unit_id=UUID(created["units"][0]["id"]), price=Decimal(10)),
    )  # type: ignore[index]
    await session.commit()
    with pytest.raises(DBAPIError):
        async with session.begin_nested():
            session.add(
                Price(
                    price_list_id=default.id,
                    product_unit_id=price.product_unit_id,
                    price=Decimal(11),
                    valid_from=utcnow(),
                    created_by=context.user.id,
                )
            )
            await session.flush()
    other = await make_org(session, context.user)
    await session.commit()
    response = await client.post(
        "/api/v1/pricing/prices",
        headers=auth(context.user, other),
        json={"price_list_id": str(default.id), "product_unit_id": str(price.product_unit_id), "price": "15"},
    )
    assert response.status_code == 404


async def test_import_invalid_file_and_upload_guards(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    job_id = await upload_job(client, context, b"not a workbook")
    await imports.process_pending()
    headers = auth(context.user, context.organization)
    assert (await client.get(f"/api/v1/imports/{job_id}", headers=headers)).json()[
        "failure_reason"
    ] == "import_file_invalid"
    for filename, data, expected in [
        ("bad.csv", b"x", "file_type_not_allowed"),
        ("huge.xlsx", b"x" * (5 * 1024 * 1024 + 1), "file_too_large"),
    ]:
        response = await client.post(
            "/api/v1/imports",
            headers=headers,
            data={"kind": "PRODUCTS"},
            files={"file": (filename, data, imports.MIME)},
        )
        assert response.status_code == 422 and response.json()["error"]["code"] == expected


def test_import_row_limit_and_formulas() -> None:
    with pytest.raises(ValueError):
        imports.read_rows(
            workbook([[str(index), "Name", None, None, None, "PCS", True] for index in range(5001)]), "PRODUCTS"
        )
    rows = imports.read_rows(workbook([["SKU", "=1+1", None, None, None, "PCS", True]]), "PRODUCTS")
    assert rows[0]["_formula"] is True


async def test_bulk_prices_rollback_all(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    created = await product(client, context)
    default = await service.default_list(session, context.organization.id)
    await session.commit()
    payload = {"price_list_id": str(default.id), "product_unit_id": created["units"][0]["id"], "price": "10"}  # type: ignore[index]
    from app.core.time import new_id

    response = await client.post(
        "/api/v1/pricing/prices/bulk",
        headers=auth(context.user, context.organization),
        json={"rows": [payload, {**payload, "product_unit_id": str(new_id())}]},
    )
    assert response.status_code == 404
    assert await session.scalar(select(func.count()).select_from(Price)) == 0


@pytest.mark.parametrize(
    "role,expected", [("OWNER", 201), ("MANAGER", 201), ("OPERATOR", 403), ("WAREHOUSE", 403), ("COURIER", 403)]
)
async def test_catalog_write_permissions(client: AsyncClient, session: AsyncSession, role: str, expected: int) -> None:
    context = await setup(session)
    user = context.user
    if role != "OWNER":
        user = await make_user(session)
        await add_member(session, context.organization, user, role)
        await session.commit()
    result = await client.post(
        "/api/v1/catalog/products",
        headers=auth(user, context.organization),
        json={"sku": "SKU", "name": "Test", "base_unit": "PCS"},
    )
    assert result.status_code == expected


async def test_product_image_upload_permissions_and_attachment(client: AsyncClient, session: AsyncSession) -> None:
    context = await setup(session)
    from PIL import Image

    image = BytesIO()
    Image.new("RGB", (2, 2)).save(image, "PNG")
    payload = {"file": ("product.png", image.getvalue(), "image/png")}
    result = await client.post(
        "/api/v1/files",
        headers=auth(context.user, context.organization),
        data={"category": "PRODUCT_IMAGE"},
        files=payload,
    )
    assert result.status_code == 201, result.text
    created = await product(client, context, image_file_id=result.json()["id"])
    assert created["image_file_id"] == result.json()["id"]
    operator = await make_user(session)
    await add_member(session, context.organization, operator, "OPERATOR")
    await session.commit()
    assert (
        await client.post(
            "/api/v1/files",
            headers=auth(operator, context.organization),
            data={"category": "PRODUCT_IMAGE"},
            files=payload,
        )
    ).status_code == 403
    sub = (await session.scalars(select(Subscription))).one()
    sub.status = "FULL_BLOCK"
    await session.commit()
    assert (
        await client.post(
            "/api/v1/files",
            headers=auth(context.user, context.organization),
            data={"category": "PRODUCT_IMAGE"},
            files=payload,
        )
    ).status_code == 403
