"""P12 §2 exports: the request, the worker, the file format and the download rules."""

import csv
from datetime import timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from uuid import uuid4

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.audit import AuditLog
from app.core.storage import get_storage
from app.core.time import utcnow
from app.modules.files.models import StoredFile
from app.modules.identity.deps import OrgContext
from app.modules.identity.models import Membership
from app.modules.reports import exports
from app.modules.reports.exports import CSV_SEPARATOR
from app.modules.reports.models import Export
from tests.factories import auth, make_user
from tests.test_reports import plain, window
from tests.test_returns_service import delivered, member


async def request_export(client: AsyncClient, ctx: OrgContext, kind: str = "sales_summary", fmt: str = "CSV"):
    answer = await client.post(
        "/api/v1/exports",
        headers={**auth(ctx.user, ctx.organization), "Idempotency-Key": str(uuid4())},
        json={"kind": kind, "format": fmt, "params": window()},
    )
    return answer


async def ready_export(
    client: AsyncClient, session: AsyncSession, ctx: OrgContext, kind: str = "sales_summary", fmt: str = "CSV"
) -> Export:
    created = await request_export(client, ctx, kind, fmt)
    assert created.status_code == 202, created.text
    await session.commit()
    assert await exports.process_pending() == 1
    export = await session.get(Export, created.json()["id"], populate_existing=True)
    assert export is not None and export.status == "READY", export.error if export else None
    return export


async def download_bytes(session: AsyncSession, export: Export) -> bytes:
    stored = await session.get(StoredFile, export.file_id)
    assert stored is not None
    return await get_storage().read(stored.storage_key)


async def test_exp_001_create_is_accepted_and_the_worker_builds_the_file(client: AsyncClient, session: AsyncSession):
    ctx, _store_ctx, order, _item = await delivered(client, session)
    await session.commit()

    created = await request_export(client, ctx)
    assert created.status_code == 202, created.text
    body = created.json()
    assert body["status"] == "PENDING" and body["kind"] == "sales_summary"

    await session.commit()
    assert await exports.process_pending() == 1
    export = await session.get(Export, body["id"], populate_existing=True)
    assert export is not None
    assert export.status == "READY" and export.row_count == 1
    assert export.ready_at is not None and export.expires_at == export.ready_at + exports.FILE_TTL

    state = await client.get(f"/api/v1/exports/{export.id}", headers=plain(ctx))
    assert state.status_code == 200 and state.json()["row_count"] == 1
    listed = await client.get("/api/v1/exports", headers=plain(ctx), params={"status": "READY"})
    assert listed.status_code == 200 and listed.json()["count"] == 1
    assert Decimal(order.total or 0) > 0


async def test_exp_006_csv_format_bom_separator(client: AsyncClient, session: AsyncSession):
    """UTF-8 with a BOM, `;` as the separator, money with a dot and headers in the requester's language."""
    ctx, _store_ctx, order, _item = await delivered(client, session)
    ctx.user.language = "ru"
    await session.commit()

    export = await ready_export(client, session, ctx)
    raw = await download_bytes(session, export)
    assert raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode("utf-8-sig")
    rows = list(csv.reader(StringIO(text), delimiter=CSV_SEPARATOR))
    assert rows[0] == ["Период", "Заказы", "Итого", "Скидка", "Стоимость доставки"]
    assert rows[1][2] == f"{order.total:.2f}" and "," not in rows[1][2]


async def test_exp_006_xlsx_money_is_a_number(client: AsyncClient, session: AsyncSession):
    ctx, _store_ctx, order, _item = await delivered(client, session)
    await session.commit()

    export = await ready_export(client, session, ctx, fmt="XLSX")
    book = load_workbook(BytesIO(await download_bytes(session, export)))
    sheet = book.active
    assert sheet is not None
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header[0] == "Давра" and header[2] == "Ҷамъ"
    money = list(sheet.iter_rows(min_row=2, max_row=2))[0][2]
    assert money.number_format == "#,##0.00"
    assert Decimal(str(money.value)) == order.total


async def test_exp_007_headers_follow_the_requesting_user(client: AsyncClient, session: AsyncSession):
    ctx, _store_ctx, _order, _item = await delivered(client, session)
    manager = await member(session, ctx, "MANAGER")
    manager.user.language = "en"
    await session.commit()

    export = await ready_export(client, session, manager)
    text = (await download_bytes(session, export)).decode("utf-8-sig")
    assert text.splitlines()[0].split(CSV_SEPARATOR)[0] == "Period"


async def test_exp_003_download_is_signed_for_fifteen_minutes_and_audited(client: AsyncClient, session: AsyncSession):
    ctx, _store_ctx, _order, _item = await delivered(client, session)
    await session.commit()
    export = await ready_export(client, session, ctx)

    answer = await client.get(f"/api/v1/exports/{export.id}/download", headers=plain(ctx))
    assert answer.status_code == 200, answer.text
    body = answer.json()
    assert answer.headers["cache-control"] == "no-store"
    expires = utcnow() + exports.DOWNLOAD_TTL
    from datetime import datetime

    signed_until = datetime.fromisoformat(body["expires_at"])
    assert timedelta(seconds=0) <= expires - signed_until < timedelta(seconds=30)
    # The signed URL is the credential: it carries no bearer token and works on its own.
    content = await client.get(body["url"])
    assert content.status_code == 200 and content.content.startswith(b"\xef\xbb\xbf")

    recorded = await session.scalars(
        select(AuditLog).where(AuditLog.action == "export.downloaded", AuditLog.entity_id == export.id)
    )
    assert len(list(recorded)) == 1


async def test_exp_002_revoked_member_cannot_download(client: AsyncClient, session: AsyncSession):
    """The access is re-checked at download time: a revoked membership never reaches the file.

    IAM-010 answers a revoked membership at the request-context layer (`membership_inactive`), before the export
    is looked up at all, so the file is unreachable one step earlier than EXP-002 describes.
    """
    ctx, store_ctx, _order, _item = await delivered(client, session)
    manager = await member(session, ctx, "MANAGER")
    await session.commit()
    export = await ready_export(client, session, manager)

    membership = await session.get(Membership, manager.membership.id)
    assert membership is not None
    membership.status = "REVOKED"
    membership.revoked_at = utcnow()
    await session.commit()

    answer = await client.get(f"/api/v1/exports/{export.id}/download", headers=plain(manager))
    assert answer.status_code == 403, answer.text
    assert answer.json()["error"]["code"] == "membership_inactive"
    # A member of another organization does not learn that the export exists at all (SEC-007).
    other = await client.get(f"/api/v1/exports/{export.id}/download", headers=plain(store_ctx))
    assert other.status_code == 404


async def test_exp_002_permission_loss_hides_the_file(client: AsyncClient, session: AsyncSession):
    ctx, _store_ctx, _order, _item = await delivered(client, session)
    manager = await member(session, ctx, "MANAGER")
    await session.commit()
    export = await ready_export(client, session, manager)

    membership = await session.get(Membership, manager.membership.id)
    assert membership is not None
    membership.role = "OPERATOR"  # no reports.sales any more (§1.2)
    await session.commit()

    answer = await client.get(f"/api/v1/exports/{export.id}/download", headers=plain(manager))
    assert answer.status_code == 403 and answer.json()["error"]["code"] == "permission_denied"


async def test_exp_004_expired_410(client: AsyncClient, session: AsyncSession):
    ctx, _store_ctx, _order, _item = await delivered(client, session)
    await session.commit()
    export = await ready_export(client, session, ctx)
    stored = await session.get(StoredFile, export.file_id)
    assert stored is not None
    key = stored.storage_key

    export.expires_at = utcnow() - timedelta(seconds=1)
    await session.commit()
    assert await exports.expire_files(session) == 1
    await session.commit()

    refreshed = await session.get(Export, export.id, populate_existing=True)
    assert refreshed is not None and refreshed.status == "EXPIRED" and refreshed.file_id is None
    answer = await client.get(f"/api/v1/exports/{export.id}/download", headers=plain(ctx))
    assert answer.status_code == 410 and answer.json()["error"]["code"] == "export_expired"
    gone = await client.get(f"/api/v1/files/content/{key}", params={"expires": 0, "signature": "x"})
    assert gone.status_code == 403


async def test_export_not_ready_is_409(client: AsyncClient, session: AsyncSession):
    ctx, _store_ctx, _order, _item = await delivered(client, session)
    await session.commit()
    created = await request_export(client, ctx)
    assert created.status_code == 202
    answer = await client.get(f"/api/v1/exports/{created.json()['id']}/download", headers=plain(ctx))
    assert answer.status_code == 409 and answer.json()["error"]["code"] == "export_not_ready"


async def test_export_unauthorized_kind_403(client: AsyncClient, session: AsyncSession):
    """EXP-001: the permission of the data, not of exporting in general."""
    ctx, store_ctx, _order, _item = await delivered(client, session)
    operator = await member(session, ctx, "OPERATOR")
    await session.commit()

    warehouse = await member(session, ctx, "WAREHOUSE")

    denied = await request_export(client, operator, "sales_summary")
    assert denied.status_code == 403 and denied.json()["error"]["code"] == "permission_denied"
    # A warehouse role reads stock but no money, so the statement is not theirs to export.
    no_money = await request_export(client, warehouse, "statement")
    assert no_money.status_code == 403
    admin_only = await request_export(client, ctx, "audit_logs")
    assert admin_only.status_code == 403
    wrong_side = await request_export(client, store_ctx, "products")
    assert wrong_side.status_code == 404
    unknown = await request_export(client, ctx, "everything")
    assert unknown.status_code == 404
    allowed = await request_export(client, operator, "orders")
    assert allowed.status_code == 202, allowed.text


async def test_exports_are_scoped_to_their_organization(client: AsyncClient, session: AsyncSession):
    ctx, store_ctx, _order, _item = await delivered(client, session)
    await session.commit()
    export = await ready_export(client, session, ctx)

    assert (await client.get(f"/api/v1/exports/{export.id}", headers=plain(store_ctx))).status_code == 404
    listed = await client.get("/api/v1/exports", headers=plain(store_ctx))
    assert listed.status_code == 200 and listed.json()["count"] == 0


async def test_data_exports_carry_their_own_columns(client: AsyncClient, session: AsyncSession):
    ctx, _store_ctx, order, _item = await delivered(client, session)
    await session.commit()

    export = await ready_export(client, session, ctx, kind="orders")
    text = (await download_bytes(session, export)).decode("utf-8-sig")
    lines = text.splitlines()
    assert lines[0].split(CSV_SEPARATOR)[0] == "Давра" or lines[0].startswith("Рақами фармоиш")
    assert order.order_number in text

    statement_export = await client.post(
        "/api/v1/exports",
        headers={**auth(ctx.user, ctx.organization), "Idempotency-Key": str(uuid4())},
        json={"kind": "statement", "format": "CSV", "params": {}},
    )
    assert statement_export.status_code == 422
    assert statement_export.json()["error"]["details"]["field"] == "partnership_id"


async def test_exp_008_rate_limit_stops_the_twenty_first_request(client: AsyncClient, session: AsyncSession):
    ctx, _store_ctx, _order, _item = await delivered(client, session)
    await session.commit()
    for _ in range(20):
        assert (await request_export(client, ctx)).status_code == 202
    limited = await request_export(client, ctx)
    assert limited.status_code == 429 and limited.json()["error"]["code"] == "rate_limited"


async def test_export_ready_notifies_the_requester(client: AsyncClient, session: AsyncSession):
    from app.core.outbox import OutboxEvent, dispatch_pending
    from app.modules.notifications.models import Notification

    ctx, _store_ctx, _order, _item = await delivered(client, session)
    await session.commit()
    export = await ready_export(client, session, ctx)
    await session.commit()

    event = await session.scalar(select(OutboxEvent).where(OutboxEvent.event_type == "EXPORT_READY"))
    assert event is not None and str(export.id) == event.payload["export_id"]
    await dispatch_pending()
    notification = await session.scalar(select(Notification).where(Notification.event_type == "EXPORT_READY"))
    assert notification is not None and notification.user_id == ctx.user.id


async def test_worker_fails_the_row_without_losing_it(client: AsyncClient, session: AsyncSession, monkeypatch):
    ctx, _store_ctx, _order, _item = await delivered(client, session)
    await session.commit()
    created = await request_export(client, ctx)
    assert created.status_code == 202
    await session.commit()

    def explode(*_args, **_kwargs):
        raise OSError("storage down")

    monkeypatch.setattr("app.modules.reports.exports.get_storage", explode)
    assert await exports.process_pending() == 1
    export = await session.get(Export, created.json()["id"], populate_existing=True)
    assert export is not None and export.status == "FAILED" and export.error == "OSError"
    state = await client.get(f"/api/v1/exports/{export.id}", headers=plain(ctx))
    assert state.json()["error"] == "OSError"


async def test_exp_005_row_limit_fails_the_export(client: AsyncClient, session: AsyncSession, monkeypatch):
    ctx, _store_ctx, _order, _item = await delivered(client, session)
    await session.commit()
    created = await request_export(client, ctx)
    assert created.status_code == 202
    await session.commit()

    monkeypatch.setattr("app.modules.reports.exports.MAX_ROWS", 0)
    assert await exports.process_pending() == 1
    export = await session.get(Export, created.json()["id"], populate_existing=True)
    assert export is not None and export.status == "FAILED" and export.error == "export_row_limit"


async def test_admin_export_belongs_to_the_platform(client: AsyncClient, session: AsyncSession):
    """ADM-006: an audit export has no organization and only its requester may download it."""
    admin = await make_user(session, is_superadmin=True)
    other = await make_user(session, is_superadmin=True)
    await session.commit()

    created = await client.post(
        "/api/v1/admin/audit-logs/export",
        headers={"Authorization": f"Bearer {auth(admin)['Authorization'].split()[1]}", "Idempotency-Key": str(uuid4())},
        params={"format": "CSV"},
    )
    assert created.status_code == 202, created.text
    await session.commit()
    assert await exports.process_pending() == 1
    export = await session.get(Export, created.json()["id"], populate_existing=True)
    assert export is not None and export.status == "READY" and export.organization_id is None

    link = await client.get(f"/api/v1/admin/exports/{export.id}/download", headers=auth(admin))
    assert link.status_code == 200, link.text
    stranger = await client.get(f"/api/v1/admin/exports/{export.id}/download", headers=auth(other))
    assert stranger.status_code == 404


async def test_exp_005_a_lost_run_is_failed_not_left_running(client: AsyncClient, session: AsyncSession):
    """A worker that disappears mid-run must not leave the request stuck in RUNNING."""
    ctx, _store_ctx, _order, _item = await delivered(client, session)
    await session.commit()
    created = await request_export(client, ctx)
    assert created.status_code == 202
    await session.commit()

    export = await session.get(Export, created.json()["id"], populate_existing=True)
    assert export is not None
    export.status = "RUNNING"
    export.started_at = utcnow() - exports.MAX_RUNTIME - timedelta(seconds=1)
    await session.commit()

    assert await exports.process_pending() == 0
    refreshed = await session.get(Export, export.id, populate_existing=True)
    assert refreshed is not None and refreshed.status == "FAILED" and refreshed.error == "export_timeout"
