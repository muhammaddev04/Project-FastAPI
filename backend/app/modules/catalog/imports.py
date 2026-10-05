from __future__ import annotations

import asyncio
import hashlib
from io import BytesIO
from typing import Any, Literal
from uuid import UUID
from zipfile import ZipFile

from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import audit
from app.core.db import get_sessionmaker
from app.core.errors import AppError
from app.core.events import DomainEvent, event_bus
from app.core.storage import get_storage, object_key
from app.core.time import utcnow
from app.modules.catalog import service
from app.modules.catalog.models import (
    Category,
    ImportError,
    ImportJob,
    ImportRow,
    Price,
    PriceList,
    Product,
    ProductUnit,
)
from app.modules.catalog.schemas import CategoryIn, PriceIn, ProductIn, ProductPatch
from app.modules.files.models import StoredFile
from app.modules.files.service import display_name
from app.modules.identity.deps import OrgContext
from app.modules.identity.models import Membership, Organization, User

Kind = Literal["PRODUCTS", "PRICES"]
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS: dict[str, dict[str, list[str]]] = {
    "PRODUCTS": {
        "en": ["sku", "name", "category", "barcode", "description", "base_unit", "is_active"],
        "ru": ["Артикул", "Название", "Категория", "Штрихкод", "Описание", "Базовая единица", "Активен"],
        "tg": ["Артикул", "Ном", "Категория", "Штрихкод", "Тавсиф", "Воҳиди асосӣ", "Фаъол"],
    },
    "PRICES": {
        "en": ["price_list_code", "sku", "unit_code", "price", "valid_from"],
        "ru": ["Код прайс-листа", "Артикул", "Код единицы", "Цена", "Действует с"],
        "tg": ["Коди рӯйхати нарх", "Артикул", "Коди воҳид", "Нарх", "Аз сана"],
    },
}
INSTRUCTIONS = {
    "en": (
        "Up to 5000 rows. SKU is unique. Base units: PCS, KG, G, L, ML, M, PACK. "
        "Prices > 0; optional dates use ISO 8601 with timezone. Confirm applies every row or none."
    ),
    "ru": (
        "До 5000 строк. Артикул уникален. Единицы: PCS, KG, G, L, ML, M, PACK. "
        "Цена > 0; даты ISO 8601 с часовым поясом. Применяются все строки или ни одна."
    ),
    "tg": (
        "То 5000 сатр. Артикул ягона аст. Воҳидҳо: PCS, KG, G, L, ML, M, PACK. "
        "Нарх > 0; сана ISO 8601 бо минтақаи вақт. Ҳама сатрҳо ё ҳеҷ кадом татбиқ мешаванд."
    ),
}


def template(kind: Kind, language: str) -> bytes:
    language = language if language in INSTRUCTIONS else "tg"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = kind
    sheet.append(HEADERS[kind][language])
    sheet.freeze_panes = "A2"
    for index in range(1, len(HEADERS[kind][language]) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 25
    book.create_sheet({"en": "Instructions", "ru": "Инструкция", "tg": "Дастур"}[language]).append(
        [INSTRUCTIONS[language]]
    )
    data = BytesIO()
    book.save(data)
    return data.getvalue()


async def upload(session: AsyncSession, context: OrgContext, file: UploadFile, kind: Kind) -> ImportJob:
    await service.writable(session, context)
    if kind == "PRICES":
        await service.default_list(session, context.organization.id)
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise AppError("file_type_not_allowed", 422)
    data = await file.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024:
        raise AppError("file_too_large", 422, {"max_bytes": 5 * 1024 * 1024})
    if not data:
        raise AppError("import_file_invalid", 422)
    key = object_key(context.organization.id, "IMPORT", "xlsx")
    await get_storage().put_private(key, data, MIME)
    stored = StoredFile(
        organization_id=context.organization.id,
        category="IMPORT",
        storage_key=key,
        content_type=MIME,
        size_bytes=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
        display_name=display_name(file.filename, "xlsx"),
        uploaded_by=context.user.id,
    )
    session.add(stored)
    await session.flush()
    job = ImportJob(company_id=context.organization.id, kind=kind, file_id=stored.id, created_by=context.user.id)
    session.add(job)
    await session.flush()
    await audit.record(
        session, "import.uploaded", "import", job.id, actor_id=context.user.id, org_id=context.organization.id
    )
    # Beat polls the durable UPLOADED state after commit; retries cannot race an uncommitted row.
    return job


async def owned(session: AsyncSession, context: OrgContext, job_id: UUID, *, lock: bool = False) -> ImportJob:
    query = select(ImportJob).where(ImportJob.company_id == context.organization.id, ImportJob.id == job_id)
    if lock:
        query = query.with_for_update().execution_options(populate_existing=True)
    job = await session.scalar(query)
    if job is None:
        raise AppError("not_found", 404)
    return job


async def transition(session: AsyncSession, job: ImportJob, target: str, *, actor_id: UUID | None = None) -> None:
    allowed = {
        "UPLOADED": {"VALIDATED", "FAILED", "CANCELLED"},
        "VALIDATED": {"CONFIRMED", "CANCELLED"},
        "CONFIRMED": {"COMPLETED", "FAILED"},
    }
    if target not in allowed.get(job.status, set()):
        raise AppError("invalid_state_transition", 409, {"from": job.status, "to": target})
    job.status = target
    job.version += 1
    await audit.record(session, f"import.{target.lower()}", "import", job.id, actor_id=actor_id, org_id=job.company_id)
    if target in {"COMPLETED", "FAILED"}:
        await event_bus.publish(
            session,
            DomainEvent(
                f"IMPORT_{target}",
                {
                    "import_id": str(job.id),
                    "company_id": str(job.company_id),
                    "created_by": str(job.created_by),
                    "summary": job.summary,
                },
                org_id=job.company_id,
            ),
        )


def read_rows(data: bytes, kind: str) -> list[dict[str, Any]]:
    with ZipFile(BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
            raise ValueError("Oversized workbook")
    book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        sheet = book.worksheets[0]
        cells = iter(sheet.iter_rows())
        raw_headers = [str(cell.value or "").strip() for cell in next(cells)]
        canonical = HEADERS[kind]["en"]
        aliases = {
            label: key for labels in HEADERS[kind].values() for key, label in zip(canonical, labels, strict=True)
        }
        headers = [aliases.get(label, label) for label in raw_headers]
        required = (
            {"sku", "name", "base_unit"} if kind == "PRODUCTS" else {"price_list_code", "sku", "unit_code", "price"}
        )
        if (
            not required <= set(headers)
            or len(set(headers)) != len(headers)
            or any(header not in canonical for header in headers)
        ):
            raise ValueError("Invalid headers")
        result: list[dict[str, Any]] = []
        for number, row in enumerate(cells, 2):
            if all(cell.value is None for cell in row):
                continue
            if len(result) >= 5000 or number > 5002:
                raise ValueError("Too many rows")
            values: dict[str, Any] = {
                header: str(cell.value).strip() if cell.value is not None else None
                for header, cell in zip(headers, row, strict=False)
            }
            values["_row"] = number
            values["_formula"] = any(cell.data_type == "f" for cell in row)
            result.append(values)
        if not result:
            raise ValueError("Empty workbook")
        return result
    finally:
        book.close()


def product_payload(data: dict[str, Any], category_id: UUID | None = None) -> ProductIn:
    active = str(data.get("is_active") or "true").lower()
    if active not in {"true", "false", "1", "0", "yes", "no"}:
        raise ValueError("is_active")
    return ProductIn(
        sku=data["sku"],
        name=data["name"],
        base_unit=data["base_unit"],
        category_id=category_id,
        barcode=data.get("barcode"),
        description=data.get("description"),
        is_active=active in {"true", "1", "yes"},
    )


async def price_payload(session: AsyncSession, company_id: UUID, data: dict[str, Any]) -> PriceIn:
    price_list = await session.scalar(
        select(PriceList).where(
            PriceList.company_id == company_id,
            PriceList.code == data.get("price_list_code"),
            PriceList.is_active.is_(True),
        )
    )
    unit = await session.scalar(
        select(ProductUnit)
        .join(Product)
        .where(
            Product.company_id == company_id,
            Product.sku == data.get("sku"),
            ProductUnit.code == data.get("unit_code"),
            ProductUnit.is_active.is_(True),
        )
    )
    if price_list is None:
        raise AppError("not_found", 422, {"field": "price_list_code"})
    if unit is None:
        raise AppError("not_found", 422, {"field": "unit_code"})
    payload = PriceIn.model_validate(
        {
            "price_list_id": price_list.id,
            "product_unit_id": unit.id,
            "price": data.get("price"),
            "valid_from": data.get("valid_from") or None,
        }
    )
    if payload.valid_from and payload.valid_from <= utcnow():
        raise ValueError("valid_from")
    return payload


async def parse_import(session: AsyncSession, job: ImportJob) -> None:
    if job.status != "UPLOADED":
        return
    file = await session.get(StoredFile, job.file_id)
    assert file is not None
    try:
        rows = await asyncio.to_thread(read_rows, await get_storage().read(file.storage_key), job.kind)
    except Exception:
        job.failure_reason = "import_file_invalid"
        await transition(session, job, "FAILED")
        return
    seen: set[str] = set()
    barcodes: set[str] = set()
    new_categories: set[str] = set()
    errors = 0
    for data in rows:
        number = int(data.pop("_row"))
        formula = data.pop("_formula")
        action = "CREATE"
        field = "row"
        try:
            if formula:
                raise ValueError("Formulas are not allowed")
            token = (
                str(data.get("sku"))
                if job.kind == "PRODUCTS"
                else f"{data.get('price_list_code')}:{data.get('sku')}:{data.get('unit_code')}"
            )
            if token in seen:
                raise ValueError("Duplicate row")
            seen.add(token)
            if job.kind == "PRODUCTS":
                payload = product_payload(data)
                existing = await session.scalar(
                    select(Product).where(Product.company_id == job.company_id, Product.sku == payload.sku)
                )
                if existing:
                    action = "UPDATE"
                    if existing.base_unit != payload.base_unit:
                        raise AppError("base_unit_immutable", 422)
                if payload.barcode:
                    collision = await session.scalar(
                        select(Product.id).where(
                            Product.company_id == job.company_id,
                            Product.barcode == payload.barcode,
                            Product.sku != payload.sku,
                        )
                    )
                    if payload.barcode in barcodes or collision:
                        raise AppError("barcode_taken", 422)
                    barcodes.add(payload.barcode)
                if data.get("category"):
                    category = await session.scalar(
                        select(Category).where(
                            Category.company_id == job.company_id,
                            Category.parent_id.is_(None),
                            func.lower(Category.name) == str(data["category"]).lower(),
                        )
                    )
                    if category and not category.is_active:
                        raise ValueError("Inactive category")
                    if not category:
                        CategoryIn(name=data["category"])
                        data["category_will_create"] = True
                        new_categories.add(data["category"])
            else:
                price_input = await price_payload(session, job.company_id, data)
                if await session.scalar(
                    select(Price.id)
                    .where(
                        Price.price_list_id == price_input.price_list_id,
                        Price.product_unit_id == price_input.product_unit_id,
                        Price.valid_to.is_(None),
                        Price.valid_from > utcnow(),
                    )
                    .limit(1)
                ):
                    raise AppError("price_overlap", 422, {"field": "price"})
        except (ValidationError, ValueError, AppError) as exc:
            code = exc.code if isinstance(exc, AppError) else "validation_error"
            if isinstance(exc, ValidationError):
                field = str(exc.errors()[0]["loc"][0])
            elif isinstance(exc, AppError):
                field = str(
                    exc.details.get("field")
                    or {"base_unit_immutable": "base_unit", "barcode_taken": "barcode"}.get(exc.code, "row")
                )
            session.add(
                ImportError(
                    import_id=job.id,
                    row_number=number,
                    field=field,
                    error_code=code,
                    message_key=f"errors.{code}",
                    params={},
                )
            )
            errors += 1
        session.add(ImportRow(import_id=job.id, row_number=number, data=data, action=action))
    job.total_rows = len(rows)
    job.error_count = errors
    job.summary = {"new_categories": sorted(new_categories)}
    await transition(session, job, "VALIDATED")


async def confirm(session: AsyncSession, context: OrgContext, job_id: UUID) -> ImportJob:
    await service.writable(session, context)
    job = await owned(session, context, job_id, lock=True)
    if job.status != "VALIDATED":
        raise AppError("import_not_ready", 409)
    if job.error_count:
        raise AppError("import_has_errors", 422)
    job.confirmed_by = context.user.id
    await transition(session, job, "CONFIRMED", actor_id=context.user.id)
    return job


async def worker_context(session: AsyncSession, job: ImportJob) -> OrgContext:
    user = await session.get(User, job.confirmed_by or job.created_by)
    organization = await session.get(Organization, job.company_id)
    membership = await session.scalar(
        select(Membership).where(
            Membership.user_id == (job.confirmed_by or job.created_by),
            Membership.organization_id == job.company_id,
            Membership.status == "ACTIVE",
        )
    )
    if user is None or user.status != "ACTIVE" or organization is None or membership is None:
        raise AppError("permission_denied", 403)
    context = OrgContext(user, membership, organization)
    if "import.run" not in context.permissions:
        raise AppError("permission_denied", 403)
    return context


async def run_import(session: AsyncSession, job: ImportJob) -> None:
    if job.status != "CONFIRMED":
        return
    context = await worker_context(session, job)
    await service.writable(session, context)
    summary = {"created": 0, "updated": 0, "skipped": 0}
    rows = list(
        await session.scalars(select(ImportRow).where(ImportRow.import_id == job.id).order_by(ImportRow.row_number))
    )
    for row in rows:
        data = row.data
        if job.kind == "PRODUCTS":
            category = None
            if data.get("category"):
                category = await session.scalar(
                    select(Category).where(
                        Category.company_id == job.company_id,
                        Category.parent_id.is_(None),
                        func.lower(Category.name) == str(data["category"]).lower(),
                    )
                )
                if category is None:
                    category = await service.create_category(session, context, CategoryIn(name=data["category"]))
            payload = product_payload(data, category.id if category else None)
            product = await session.scalar(
                select(Product).where(Product.company_id == job.company_id, Product.sku == payload.sku)
            )
            if product:
                if payload.base_unit != product.base_unit:
                    raise AppError("base_unit_immutable", 409)
                await service.update_product(
                    session,
                    context,
                    product.id,
                    ProductPatch(version=product.version, **payload.model_dump(exclude={"base_unit", "image_file_id"})),
                )
                summary["updated"] += 1
            else:
                await service.create_product(session, context, payload)
                summary["created"] += 1
        else:
            await service.set_price(session, context, await price_payload(session, job.company_id, data))
            summary["created"] += 1
    job.summary = summary
    await transition(session, job, "COMPLETED", actor_id=context.user.id)


async def process_pending(limit: int = 20) -> int:
    processed = 0
    for _ in range(limit):
        async with get_sessionmaker()() as session, session.begin():
            candidate = await session.scalar(
                select(ImportJob)
                .join(Organization, Organization.id == ImportJob.company_id)
                .where(ImportJob.status.in_(["UPLOADED", "CONFIRMED"]))
                .order_by(ImportJob.created_at, ImportJob.id)
                .with_for_update(of=Organization, skip_locked=True)
                .limit(1)
            )
            if candidate is None:
                break
            # Match API lock order (organization before job), including cancel/confirm races.
            job = await session.scalar(
                select(ImportJob)
                .where(ImportJob.id == candidate.id, ImportJob.status.in_(["UPLOADED", "CONFIRMED"]))
                .with_for_update(skip_locked=True)
                .execution_options(populate_existing=True)
            )
            if job is None:
                break
            try:
                # Job lock is outside the savepoint: failure rolls back EVERY business mutation.
                async with session.begin_nested():
                    if job.status == "UPLOADED":
                        await parse_import(session, job)
                    else:
                        await run_import(session, job)
                    await session.flush()
            except Exception as exc:
                await session.refresh(job)
                job.failure_reason = exc.code if isinstance(exc, AppError) else "import_failed"
                await transition(session, job, "FAILED")
            processed += 1
    return processed
